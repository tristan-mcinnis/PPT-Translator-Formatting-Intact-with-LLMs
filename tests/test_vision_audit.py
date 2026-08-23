"""Tests for the vision audit module."""
from __future__ import annotations

import base64
import json
from pathlib import Path

import pytest

from ppt_translator.vision_audit import (
    DEFAULT_VISION_MODEL,
    VisionAuditError,
    VisionAuditor,
    collect_slide_texts,
    parse_vision_findings,
    print_audit_summary,
    render_deck_to_images,
    write_vision_audit_report,
)


class _FakeMessage:
    def __init__(self, content):
        self.content = content


class _FakeChoice:
    def __init__(self, content):
        self.message = _FakeMessage(content)


class _FakeResponse:
    def __init__(self, content):
        self.choices = [_FakeChoice(content)]


class _FakeCompletions:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return _FakeResponse(self.responses.pop(0))


class _FakeChat:
    def __init__(self, completions):
        self.completions = completions


class _FakeClient:
    def __init__(self, responses):
        self.chat = _FakeChat(_FakeCompletions(responses))


def _png_file(tmp_path, name="slide-1.png", payload=b"fake-png-bytes"):
    path = tmp_path / name
    path.write_bytes(payload)
    return path


def test_render_raises_clear_error_when_tools_missing(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "ppt_translator.vision_audit.shutil.which", lambda name: None
    )
    with pytest.raises(VisionAuditError) as excinfo:
        render_deck_to_images("whatever.pptx", tmp_path / "out")
    assert "soffice" in str(excinfo.value) and "pdftoppm" in str(excinfo.value)


def test_parse_vision_findings_handles_fences_and_prose():
    fenced = 'Here you go:\n```json\n{"issues": [{"slide": 1, "severity": "high", "type": "truncated", "description": "cut"}]}\n```'
    issues = parse_vision_findings(fenced)
    assert len(issues) == 1 and issues[0]["type"] == "truncated"

    plain = 'Sure. {"issues": []} hope that helps'
    assert parse_vision_findings(plain) == []

    assert parse_vision_findings(None) == []
    assert parse_vision_findings("no json at all") == []


def test_audit_slide_sends_base64_image_and_parses_issues(tmp_path):
    image = _png_file(tmp_path)
    response = json.dumps(
        {
            "issues": [
                {"slide": 1, "severity": "HIGH", "type": "Text_Overflow", "description": "spills"},
                {"slide": 1, "severity": "bogus", "description": ""},
            ]
        }
    )
    client = _FakeClient([response])
    auditor = VisionAuditor(client)

    issues = auditor.audit_slide(image, 1, ["你好"])

    call = client.chat.completions.calls[0]
    assert call["model"] == DEFAULT_VISION_MODEL
    user_content = call["messages"][1]["content"]
    assert user_content[0]["type"] == "text"
    assert "你好" in user_content[0]["text"]
    assert user_content[1]["type"] == "image_url"
    encoded = base64.b64encode(b"fake-png-bytes").decode()
    assert user_content[1]["image_url"]["url"].endswith(encoded)

    assert issues[0] == {
        "slide": 1,
        "severity": "high",
        "type": "text_overflow",
        "description": "spills",
    }
    # Malformed second issue is normalised rather than dropped.
    assert issues[1]["severity"] == "low"
    assert issues[1]["type"] == "other"


def test_audit_deck_collects_and_sorts_issues(tmp_path):
    images = [
        _png_file(tmp_path, "slide-2.png"),
        _png_file(tmp_path, "slide-1.png"),
    ]
    responses = ['{"issues": []}', '{"issues": [{"slide": 2, "severity": "high", "description": "bad"}]}']
    auditor = VisionAuditor(_FakeClient(responses))

    issues = auditor.audit_deck(images, expected_texts_by_slide={2: ["文本"]})

    assert [i["slide"] for i in issues] == [2]
    assert issues[0]["severity"] == "high"


def test_audit_deck_reports_per_slide_errors_without_aborting(tmp_path):
    class ExplodingClient:
        class chat:
            class completions:
                @staticmethod
                def create(**kwargs):
                    raise RuntimeError("boom")

    images = [_png_file(tmp_path, "slide-1.png")]
    auditor = VisionAuditor(ExplodingClient())

    issues = auditor.audit_deck(images)

    assert len(issues) == 1
    assert issues[0]["type"] == "audit_error"
    assert "boom" in issues[0]["description"]


def test_report_writer_summarises_counts_and_groups_by_slide(tmp_path):
    report = write_vision_audit_report(
        tmp_path / "audit.md",
        "deck.pptx",
        [
            {"slide": 2, "severity": "high", "type": "truncated", "description": "clipped"},
            {"slide": 2, "severity": "low", "type": "contrast", "description": "dim"},
            {"slide": 3, "severity": "medium", "type": "untranslated", "description": "中文 remains"},
        ],
    )

    text = Path(report).read_text(encoding="utf-8")
    assert "deck.pptx" in text
    assert "**3**" in text
    assert "(high: 1, medium: 1, low: 1)" in text
    assert "## Slide 2" in text and "## Slide 3" in text
    assert "[HIGH] truncated" in text


def test_report_writer_clean_deck(tmp_path):
    report = write_vision_audit_report(tmp_path / "audit.md", "deck.pptx", [])
    text = Path(report).read_text(encoding="utf-8")
    assert "No visual issues" in text and "**0**" in text


def test_print_audit_summary_outputs_high_issues(capsys):
    print_audit_summary(
        [{"slide": 4, "severity": "high", "type": "garbled", "description": "tofu boxes"}]
    )
    out = capsys.readouterr().out
    assert "1 issue(s)" in out and "HIGH slide 4" in out

    print_audit_summary([])
    assert "passed" in capsys.readouterr().out


def test_build_default_auditor_requires_api_key(monkeypatch):
    from ppt_translator.vision_audit import build_default_vision_auditor

    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    with pytest.raises(VisionAuditError):
        build_default_vision_auditor()

    monkeypatch.setenv("DEEPSEEK_API_KEY", "k")
    monkeypatch.setenv("DEEPSEEK_API_BASE", "https://example.test")
    auditor = build_default_vision_auditor(model="custom-vision")
    assert auditor.model == "custom-vision"


def test_collect_slide_texts_includes_tables_and_textboxes(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "你好"
    rows, cols = 2, 2
    table_shape = slide.shapes.add_table(rows, cols, Inches(5), Inches(1), Inches(4), Inches(2))
    table_shape.table.cell(0, 0).text = "单元格"
    deck = tmp_path / "deck.pptx"
    prs.save(deck)

    texts = collect_slide_texts(deck)

    assert texts[1] == ["你好", "单元格"]
