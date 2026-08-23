"""Tests for the PowerPoint pipeline formatting round-trip."""
from __future__ import annotations

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ppt_translator.pipeline import (
    get_alignment_value,
    get_shape_properties,
    ppt_to_xml,
)


def _slide_with_aligned_textbox(prs, alignment):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = alignment
    run = paragraph.add_run()
    run.text = "你好世界"
    run.font.size = Pt(24)
    return slide


def test_extracted_alignment_round_trips_through_get_alignment_value():
    """Every extracted alignment string must resolve back to its enum value."""
    for alignment in (PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT, PP_ALIGN.JUSTIFY):
        prs = Presentation()
        slide = _slide_with_aligned_textbox(prs, alignment)
        props = get_shape_properties(slide.shapes[0])
        assert props["alignment"] is not None, f"alignment not captured for {alignment}"
        assert get_alignment_value(props["alignment"]) == alignment


def test_intermediate_slide_xml_contains_only_that_slide(tmp_path):
    """slide_<n>_<stage>.xml must hold slide n alone, not slides 1..n."""
    prs = Presentation()
    for text in ("第一页", "第二页"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = text
    pptx_path = tmp_path / "deck.pptx"
    prs.save(pptx_path)

    ppt_to_xml(str(pptx_path), translator=None, source_lang="zh", target_lang="en")

    first = (tmp_path / "slide_1_original.xml").read_text(encoding="utf-8")
    second = (tmp_path / "slide_2_original.xml").read_text(encoding="utf-8")
    assert 'number="1"' in first and 'number="2"' not in first
    assert 'number="2"' in second and 'number="1"' not in second