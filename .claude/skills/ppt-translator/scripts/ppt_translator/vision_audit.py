"""Visual audit of translated decks using a DeepSeek vision language model.

The audit renders every slide of the rebuilt deck to an image (via LibreOffice
headless + pdftoppm), then asks a vision model
(deepseek-v4-flash-vision-exp) to inspect each rendered slide for visual
regressions such as text overflow, truncation, garbled glyphs or leftover
untranslated source text.
"""
from __future__ import annotations

import base64
import json
import os
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any, Dict, List, Sequence

from openai import OpenAI

DEFAULT_VISION_MODEL = "deepseek-v4-flash-vision-exp"

_SEVERITY_ORDER = {"high": 0, "medium": 1, "low": 2}

_AUDIT_SYSTEM_PROMPT = (
    "You are a meticulous presentation QA inspector. You will be shown a "
    "rendered slide image together with the text that is SUPPOSED to appear "
    "on that slide. Compare the rendering against the expected content and "
    "report concrete visual problems only. Do not invent stylistic nitpicks."
)

_AUDIT_USER_TEMPLATE = (
    "Inspect this rendered PowerPoint slide (slide {slide_number}) for "
    "translation/layout defects.\n\n"
    "Check specifically for:\n"
    "1. text_overflow: text spilling outside its box, shape, table cell or slide edge\n"
    "2. truncated: text visibly cut off mid-word or clipped by container edges\n"
    "3. garbled: tofu boxes, mojibake, broken glyphs or wrong-script rendering\n"
    "4. untranslated: source-language text that should have been translated\n"
    "5. overlap: shapes or text blocks colliding/illegibly stacked\n"
    "6. contrast: text that is unreadable because of its background\n\n"
    "Expected translatable text on this slide (JSON):\n{expected_text}\n\n"
    "Respond with ONLY a JSON object, no markdown fences:\n"
    '{{"issues": [{{"slide": <int>, "severity": "high|medium|low", '
    '"type": "text_overflow|truncated|garbled|untranslated|overlap|contrast|other", '
    '"description": "<one concise sentence>"}}]}}'
)


class VisionAuditError(RuntimeError):
    """Raised when slide rendering or the vision audit cannot proceed."""


def render_deck_to_images(
    ppt_path,
    output_dir,
    *,
    dpi=100,
    soffice_bin="soffice",
    pdftoppm_bin="pdftoppm",
):
    """Render every slide of *ppt_path* to PNG files inside *output_dir*.

    Uses LibreOffice headless to convert the deck to PDF and pdftoppm to
    rasterise the PDF pages. Returns the sorted list of PNG paths.

    Raises VisionAuditError when either tool is unavailable or a conversion
    step fails.
    """
    ppt_path = Path(ppt_path)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    soffice = shutil.which(soffice_bin)
    pdftoppm = shutil.which(pdftoppm_bin)
    missing = []
    if not soffice:
        missing.append("'soffice' (LibreOffice, required to render slides)")
    if not pdftoppm:
        missing.append("'pdftoppm' (poppler, required to convert the PDF to images)")
    if missing:
        raise VisionAuditError(
            "Visual audit needs external tools that were not found: "
            + ", ".join(missing)
        )

    # Isolated profile avoids clashing with a running LibreOffice instance.
    profile_dir = output_dir / ".lo_profile"
    convert_cmd = [
        str(soffice),
        "-env:UserInstallation=" + profile_dir.as_uri(),
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir),
        str(ppt_path),
    ]
    try:
        result = subprocess.run(convert_cmd, capture_output=True, text=True, timeout=300)
    except subprocess.TimeoutExpired as exc:
        raise VisionAuditError(f"LibreOffice conversion timed out for {ppt_path.name}") from exc
    pdf_path = output_dir / (ppt_path.stem + ".pdf")
    if result.returncode != 0 or not pdf_path.exists():
        raise VisionAuditError(
            f"LibreOffice failed to convert {ppt_path.name} to PDF: "
            + (result.stderr or result.stdout or "no output").strip()
        )

    raster_cmd = [str(pdftoppm), "-png", "-r", str(dpi), str(pdf_path), str(output_dir / "slide")]
    try:
        result = subprocess.run(raster_cmd, capture_output=True, text=True, timeout=300)
    except subprocess.TimeoutExpired as exc:
        raise VisionAuditError(f"pdftoppm timed out for {pdf_path.name}") from exc
    if result.returncode != 0:
        raise VisionAuditError(
            f"pdftoppm failed on {pdf_path.name}: {(result.stderr or '').strip()}"
        )

    images = sorted(p for p in output_dir.glob("slide*.png") if p.is_file())
    if not images:
        raise VisionAuditError(f"No slide images were produced for {ppt_path.name}")
    return images


def parse_vision_findings(content):
    """Extract the issue list from a vision model response.

    Tolerates markdown fences or surrounding prose; returns [] when nothing
    parseable is found.
    """
    if not content:
        return []
    text = content.strip()
    fence_match = re.search(r"```(?:json)?\s*(.+?)\s*```", text, re.S)
    if fence_match:
        text = fence_match.group(1)
    decoder = json.JSONDecoder()
    for idx, ch in enumerate(text):
        if ch == "{":
            try:
                parsed, _ = decoder.raw_decode(text[idx:])
            except json.JSONDecodeError:
                continue
            if isinstance(parsed, dict) and isinstance(parsed.get("issues"), list):
                return [i for i in parsed["issues"] if isinstance(i, dict)]
    return []


def _image_data_url(image_path):
    encoded = base64.b64encode(image_path.read_bytes()).decode("ascii")
    return "data:image/png;base64," + encoded


def normalise_severity(value):
    lowered = str(value or "low").strip().lower()
    return lowered if lowered in _SEVERITY_ORDER else "low"


class VisionAuditor:
    """Audits rendered slide images through a DeepSeek vision model."""

    def __init__(self, client, *, model=DEFAULT_VISION_MODEL, temperature=0.0):
        self.client = client
        self.model = model
        self.temperature = temperature

    def build_messages(self, image_path, slide_number, expected_texts=()):
        expected_payload = json.dumps(list(expected_texts), ensure_ascii=False, indent=1)
        return [
            {"role": "system", "content": _AUDIT_SYSTEM_PROMPT},
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": _AUDIT_USER_TEMPLATE.format(
                            slide_number=slide_number,
                            expected_text=expected_payload,
                        ),
                    },
                    {
                        "type": "image_url",
                        "image_url": {"url": _image_data_url(image_path)},
                    },
                ],
            },
        ]

    def audit_slide(self, image_path, slide_number, expected_texts=()):
        """Audit one rendered slide; returns its normalised issue list."""
        messages = self.build_messages(image_path, slide_number, expected_texts)
        response = self.client.chat.completions.create(
            model=self.model,
            messages=messages,
            temperature=self.temperature,
            stream=False,
        )
        content = response.choices[0].message.content
        issues = []
        for raw in parse_vision_findings(content):
            issues.append(
                {
                    "slide": int(raw.get("slide") or slide_number),
                    "severity": normalise_severity(raw.get("severity")),
                    "type": str(raw.get("type") or "other").strip().lower() or "other",
                    "description": str(raw.get("description") or "").strip(),
                }
            )
        return issues

    def audit_deck(self, image_paths, *, expected_texts_by_slide=None):
        """Audit every rendered slide image, returning all issues found."""
        expected = expected_texts_by_slide or {}
        all_issues = []
        for index, image_path in enumerate(sorted(image_paths), start=1):
            try:
                issues = self.audit_slide(image_path, index, expected.get(index, ()))
            except Exception as exc:
                all_issues.append(
                    {
                        "slide": index,
                        "severity": "medium",
                        "type": "audit_error",
                        "description": f"Slide {index} could not be audited: {exc}",
                    }
                )
                continue
            all_issues.extend(issues)
        all_issues.sort(key=lambda i: (i["slide"], _SEVERITY_ORDER.get(i["severity"], 3)))
        return all_issues


def build_default_vision_auditor(*, model=DEFAULT_VISION_MODEL, api_key=None, base_url=None):
    """Create a VisionAuditor backed by the DeepSeek API."""
    resolved_key = api_key or os.getenv("DEEPSEEK_API_KEY")
    if not resolved_key:
        raise VisionAuditError(
            "The visual audit requires a DeepSeek API key. Set DEEPSEEK_API_KEY."
        )
    resolved_base = base_url or os.getenv("DEEPSEEK_API_BASE", "https://api.deepseek.com")
    client = OpenAI(api_key=resolved_key, base_url=resolved_base)
    return VisionAuditor(client, model=model)


def collect_slide_texts(prs_path):
    """Return the translatable text of every slide, keyed by slide number."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(prs_path))
    result = {}
    for number, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            texts.append(cell.text.strip())
            elif hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        result[number] = texts
    return result


def write_vision_audit_report(output_path, deck_name, issues, *, model=DEFAULT_VISION_MODEL):
    """Write the audit findings as a Markdown report and return its path."""
    output_path = Path(output_path)
    counts = {}
    for issue in issues:
        counts[issue["severity"]] = counts.get(issue["severity"], 0) + 1

    lines = [
        "# Visual audit report - " + deck_name,
        "",
        f"Model: `{model}` · Issues found: **{len(issues)}** "
        + f"(high: {counts.get('high', 0)}, medium: {counts.get('medium', 0)}, low: {counts.get('low', 0)})",
        "",
    ]
    if not issues:
        lines += ["No visual issues were detected.", ""]
        output_path.write_text("\n".join(lines), encoding="utf-8")
        return output_path

    current_slide = None
    for issue in issues:
        if issue["slide"] != current_slide:
            current_slide = issue["slide"]
            lines.append(f"## Slide {current_slide}")
            lines.append("")
        lines.append(
            f"- **[{issue['severity'].upper()}] {issue['type']}** - {issue['description']}"
        )
    lines.append("")
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def print_audit_summary(issues):
    """Print a compact console summary of the audit results."""
    if not issues:
        print("Visual audit passed: no issues detected.")
        return
    counts = {}
    for issue in issues:
        counts[issue["severity"]] = counts.get(issue["severity"], 0) + 1
    summary = ", ".join(f"{counts[s]} {s}" for s in ("high", "medium", "low") if s in counts)
    print(f"Visual audit found {len(issues)} issue(s): {summary}.")
    for issue in issues:
        if issue["severity"] == "high":
            print(f"  HIGH slide {issue['slide']}: {issue['description']}")
